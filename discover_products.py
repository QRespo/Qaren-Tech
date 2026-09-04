"""
سكربت اكتشاف منتجات جديدة تلقائيًا من علي إكسبريس، حسب كلمات بحث تحددها.
يضيف المنتجات مباشرة لملف products.xlsx (كعروض من علي إكسبريس)،
وبعدها تشغّل init_db.py عشان الموقع يقرأها.

طريقة الاستخدام: عدّل قائمة SEARCHES بالأسفل، وشغّل السكربت.
"""

import os
import time
import openpyxl
from dotenv import load_dotenv

load_dotenv()

ALIEXPRESS_APP_KEY = os.getenv("ALIEXPRESS_APP_KEY")
ALIEXPRESS_APP_SECRET = os.getenv("ALIEXPRESS_APP_SECRET")
ALIEXPRESS_TRACKING_ID = os.getenv("ALIEXPRESS_TRACKING_ID")

USD_TO_SAR = 3.75
MAX_RESULTS_PER_SEARCH = 15   # عدد المنتجات اللي يجيبها لكل كلمة بحث
REQUEST_DELAY_SECONDS = 5

# ============ عدّل هذي القائمة حسب اللي تبي تضيفه ============
# كل سطر: (كلمة البحث بالإنجليزي, اسم الفئة اللي تنكتب بالموقع)
SEARCHES = [
    ("gaming mouse", "ملحقات"),
    ("mechanical keyboard", "ملحقات"),
    ("cpu cooler", "مبردات"),
    ("microphone stand", "مايكروفونات"),
]
# ================================================================


def get_client():
    from aliexpress_api import AliexpressApi, models
    return AliexpressApi(
        ALIEXPRESS_APP_KEY,
        ALIEXPRESS_APP_SECRET,
        models.Language.AR,
        models.Currency.USD,
        ALIEXPRESS_TRACKING_ID,
    )


def load_existing_names(ws):
    names = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            names.add(row[0].strip())
    return names


def main():
    client = get_client()

    wb = openpyxl.load_workbook('products.xlsx')
    ws = wb['المنتجات']
    existing_names = load_existing_names(ws)

    total_added = 0

    for keyword, category in SEARCHES:
        print(f"\nجاري البحث عن: {keyword}")
        try:
            result = client.get_hotproducts(
                keywords=keyword,
                page_size=MAX_RESULTS_PER_SEARCH,
            )
        except Exception as e:
            print(f"  ❌ خطأ بالبحث: {e}")
            continue

        products = getattr(result, "products", None) or []
        if not products:
            print("  لقيت 0 منتج، جرب كلمة بحث ثانية.")
            continue

        for p in products:
            name = (p.product_title or "").strip()
            if not name or name in existing_names:
                continue

            price_usd = p.target_sale_price or p.target_original_price
            if not price_usd:
                continue
            price_sar = round(float(price_usd) * USD_TO_SAR, 2)

            image_url = p.product_main_image_url or ""
            product_url = p.promotion_link or p.product_detail_url or ""

            next_row = ws.max_row + 1
            ws.cell(row=next_row, column=1, value=name)
            ws.cell(row=next_row, column=2, value=category)
            ws.cell(row=next_row, column=3, value=image_url)
            ws.cell(row=next_row, column=4, value="علي إكسبريس")
            ws.cell(row=next_row, column=5, value=price_sar)
            ws.cell(row=next_row, column=6, value=product_url)

            existing_names.add(name)
            total_added += 1
            print(f"  ✅ أضفت: {name[:60]}... — {price_sar} ريال")

        time.sleep(REQUEST_DELAY_SECONDS)

    wb.save('products.xlsx')
    print(f"\nخلص البحث. أضفت {total_added} منتج جديد.")
    print("شغّل الآن: python init_db.py")


if __name__ == '__main__':
    main()
